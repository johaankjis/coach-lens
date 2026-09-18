"""Real-mode wiring tests use generated workbooks only, never supplied QA data."""

import asyncio
from collections import Counter
from decimal import Decimal
import json
from pathlib import Path
import subprocess
import sys

from openpyxl import load_workbook
from fastapi.testclient import TestClient
import pytest

from app.config import Settings
from app.design.demo import DemoDesignFixture
from app.design.service import DesignError, DesignService, UnavailableDesignProvider
from app.interventions.handoff import ValidatedInterventionHandoff
from app.interventions.service import (UnavailableInterventionReasoner,
                                       UnavailableSolutionValidator)
from app.diagnostics.engine import (ControlledTestReasoner, DiagnosticError, DiagnosticService,
                                    UnavailableReasoner, build_bundle, detect_signals)
from app.main import app
from app.results_cx import PipelineValidationError, normalize
from app.results_cx import demo
from app.results_cx.models import Domain
from test_results_cx import make_workbook


def generated_sources(root: Path):
    for domain, name in demo.REQUIRED_WORKBOOKS.items():
        make_workbook(root, domain.value, filename=name)


@pytest.fixture
def app_state():
    """Restore the module-level app after a test installs a different mode into it."""
    prior = (app.state.diagnostics, app.state.designs, app.state.demo_mode)
    try:
        yield app
    finally:
        app.state.diagnostics, app.state.designs, app.state.demo_mode = prior


def test_real_mode_requires_exact_workbooks_and_never_falls_back(tmp_path):
    with pytest.raises(PipelineValidationError, match="missing_demo_workbooks") as failure:
        demo.load_results_cx_demo(tmp_path / "raw")
    assert all(name in str(failure.value) for name in demo.REQUIRED_WORKBOOKS.values())
    root = tmp_path / "raw"
    generated_sources(root)
    make_workbook(root, "compliance", filename="unrelated.xlsx")
    with pytest.raises(PipelineValidationError, match="unexpected_demo_workbooks"):
        demo.load_results_cx_demo(root)


def test_one_missing_workbook_names_only_the_missing_file(tmp_path):
    root = tmp_path / "raw"
    generated_sources(root)
    absent = demo.REQUIRED_WORKBOOKS[Domain.BUSINESS_PROCESS]
    (root / absent).unlink()
    with pytest.raises(PipelineValidationError, match="missing_demo_workbooks") as failure:
        demo.load_results_cx_demo(root)
    assert absent in str(failure.value)
    assert demo.REQUIRED_WORKBOOKS[Domain.COMPLIANCE] not in str(failure.value)


def test_nested_or_miscased_copies_are_rejected_by_real_path(tmp_path):
    root = tmp_path / "raw"
    generated_sources(root)
    name = demo.REQUIRED_WORKBOOKS[Domain.COMPLIANCE]
    make_workbook(root / "archive", "compliance", filename=name)
    with pytest.raises(PipelineValidationError, match="unexpected_demo_workbooks") as failure:
        demo.load_results_cx_demo(root)
    assert f"archive/{name}" in str(failure.value)
    (root / "archive" / name).unlink()
    (root / name).rename(root / name.lower())
    # Case-insensitive filesystems satisfy is_file() for the required name; both outcomes refuse.
    with pytest.raises(PipelineValidationError) as failure:
        demo.load_results_cx_demo(root)
    assert failure.value.code in {"missing_demo_workbooks", "unexpected_demo_workbooks"}


def test_filename_must_match_m2_classified_domain(tmp_path):
    root = tmp_path / "raw"
    generated_sources(root)
    compliance = root / demo.REQUIRED_WORKBOOKS[Domain.COMPLIANCE]
    member = root / demo.REQUIRED_WORKBOOKS[Domain.MEMBER_EXPERIENCE]
    swap = root / "swap.bin"
    compliance.rename(swap)
    member.rename(compliance)
    swap.rename(member)
    with pytest.raises(PipelineValidationError, match="demo_domain_mismatch"):
        demo.load_results_cx_demo(root)


def test_malformed_and_invalid_source_rows_fail_through_m2_validation(tmp_path):
    root = tmp_path / "raw"
    generated_sources(root)
    path = root / demo.REQUIRED_WORKBOOKS[Domain.COMPLIANCE]
    workbook = load_workbook(path)
    workbook["Export"].cell(3, 7).value = "Maybe"  # Answers column of the first criterion row.
    workbook.save(path)
    with pytest.raises(PipelineValidationError, match="invalid_answer"):
        demo.load_results_cx_demo(root)
    path.write_bytes(b"not a workbook")
    with pytest.raises(PipelineValidationError, match="unreadable_workbook"):
        demo.load_results_cx_demo(root)


def test_real_mode_delegates_to_m2_and_keeps_domains_evidence_and_lineage(tmp_path, monkeypatch):
    root = tmp_path / "raw"
    generated_sources(root)
    calls = []
    original = demo.normalize

    def checked_normalize(sources):
        calls.append(set(sources))
        return original(sources)

    monkeypatch.setattr(demo, "normalize", checked_normalize)
    evaluations = demo.load_results_cx_demo(root)
    assert calls == [set(Domain)]
    assert evaluations == normalize(demo.discover_sources(root))
    assert all(e.internal_id.startswith("rcx1_") for e in evaluations)
    signals = detect_signals(evaluations)
    assert len(signals) == 6  # Two generated criteria in each of three domains.
    assert {(s.domain, s.criterion) for s in signals} == {
        (domain, question) for domain in Domain for question in ("Criterion A", "Criterion B")}
    assert [(s.fail_count, s.evaluated_results, s.fail_rate, s.feedback_count)
            for s in signals if s.criterion == "Criterion B"] == [(2, 2, Decimal(1), 2)] * 3
    # Zero-failure criteria are still observed signals, and the documented order is exact.
    assert [(s.fail_count, s.evaluated_results, s.fail_rate, s.feedback_count)
            for s in signals if s.criterion == "Criterion A"] == [(0, 2, Decimal(0), 0)] * 3
    assert [(s.fail_count, s.domain.value, s.criterion) for s in signals] == sorted(
        ((s.fail_count, s.domain.value, s.criterion) for s in signals),
        key=lambda t: (-t[0], t[1], t[2]))
    assert all(s.evaluated_evaluations == 2 and s.total_evaluations == 2 for s in signals)
    for signal in signals:
        bundle = build_bundle(signal, evaluations)
        expected = [(e.internal_id, c) for e in evaluations for c in e.criteria
                    if c.domain == signal.domain and c.question.strip() == signal.criterion]
        assert Counter((i.evaluation_id, i.source_lineage.source_filename,
                        i.source_lineage.source_sheet, i.source_lineage.excel_row)
                       for i in bundle.items) == Counter(
            (eid, c.lineage.source_filename, c.lineage.source_sheet, c.lineage.excel_row)
            for eid, c in expected)
        assert all(i.source_lineage.source_filename == demo.REQUIRED_WORKBOOKS[signal.domain]
                   for i in bundle.items)
        assert not {"cause_domain", "performance_dimension", "training_need"} & set(signal.model_dump())


def test_incomplete_criterion_coverage_and_feedback_are_counted_not_assumed(tmp_path):
    root = tmp_path / "raw"
    generated_sources(root)
    path = root / demo.REQUIRED_WORKBOOKS[Domain.COMPLIANCE]
    workbook = load_workbook(path)
    sheet = workbook["Export"]
    sheet.delete_rows(5)  # The second generated evaluation lacks Criterion B in this domain.
    sheet.cell(3, 8).value = None  # First evaluation's Criterion B has no feedback.
    workbook.save(path)
    evaluations = demo.load_results_cx_demo(root)
    signal = next(s for s in detect_signals(evaluations)
                  if s.domain == Domain.COMPLIANCE and s.criterion == "Criterion B")
    assert (signal.evaluated_evaluations, signal.evaluated_results, signal.fail_count,
            signal.fail_rate, signal.feedback_count) == (1, 1, 1, Decimal(1), 0)
    # The 100% rate covers one represented evaluation out of two loaded; both numbers are exposed.
    assert signal.total_evaluations == 2
    assert len(build_bundle(signal, evaluations).items) == 1


def test_stale_signal_from_another_dataset_is_refused(tmp_path):
    root = tmp_path / "raw"
    generated_sources(root)
    evaluations = demo.load_results_cx_demo(root)
    signal = detect_signals(evaluations)[0]
    with pytest.raises(DiagnosticError, match="Signal and criterion records disagree"):
        build_bundle(signal, evaluations + [evaluations[0].model_copy(update={"internal_id": "rcx1_other"})])


def test_real_mode_has_no_diagnostic_provider(tmp_path):
    root = tmp_path / "raw"
    generated_sources(root)
    service = DiagnosticService(demo.load_results_cx_demo(root), UnavailableReasoner())
    signal = service.list_signals()[0]
    assert service.evidence(signal.signal_id).items
    with pytest.raises(DiagnosticError, match="No diagnostic reasoning provider") as failure:
        asyncio.run(service.diagnose(signal.signal_id))
    assert failure.value.code == "reasoner_unavailable"
    assert service.list_hypotheses(signal.signal_id) == []


def test_installed_real_mode_http_surface_stops_at_diagnosis(tmp_path, monkeypatch, app_state):
    # Provider-looking configuration in the environment must be inert: nothing reads it.
    monkeypatch.setenv("AWS_ACCESS_KEY_ID", "not-a-real-key-id")
    monkeypatch.setenv("AWS_BEARER_TOKEN_BEDROCK", "not-a-real-token")
    monkeypatch.setenv("COACHLENS_API_DIAGNOSTIC_PROVIDER", "bedrock")
    assert set(Settings.model_fields) == {"title", "service_name", "diagnostic_evaluations_path",
                                          "bedrock_enabled", "bedrock_region", "bedrock_model_id"}
    root = tmp_path / "raw"
    generated_sources(root)
    demo.install_results_cx_demo(app_state, demo.load_results_cx_demo(root))
    assert isinstance(app_state.state.diagnostics.reasoner, UnavailableReasoner)
    # AWS-4: the M5 intervention step reads the intervention record; the reasoner behind it and
    # the training designer are both unavailable here.
    assert isinstance(app_state.state.designs.intervention, ValidatedInterventionHandoff)
    assert isinstance(app_state.state.interventions.reasoner, UnavailableInterventionReasoner)
    assert isinstance(app_state.state.interventions.validator, UnavailableSolutionValidator)
    assert app_state.state.interventions.diagnostics is app_state.state.diagnostics
    assert isinstance(app_state.state.designs.training, UnavailableDesignProvider)
    client = TestClient(app_state)
    assert client.get("/diagnostics/mode").json() == {
        "mode": "real_results_cx", "diagnostic_provider": "unavailable", "remote_diagnosis": "unavailable",
        "design_provider": "unavailable", "intervention_provider": "unavailable",
        "solution_validator": "unavailable", "evaluation_count": 2, "signal_count": 6}
    signals = client.get("/diagnostics/signals").json()
    assert [s["signal_id"] for s in signals] == [
        s.signal_id for s in app_state.state.diagnostics.list_signals()]
    assert all(s["total_evaluations"] == 2 for s in signals)
    review = client.get(f"/diagnostics/signals/{signals[0]['signal_id']}/review-evidence").json()
    assert review["items"] and all(item["evaluation_id"].startswith("rcx1_") for item in review["items"])
    assert all(item["source_lineage"]["source_filename"] in demo.REQUIRED_WORKBOOKS.values()
               for item in review["items"])
    provider_view = client.get(f"/diagnostics/signals/{signals[0]['signal_id']}/evidence").json()
    assert all("source_lineage" not in item for item in provider_view["items"])
    response = client.post(f"/diagnostics/signals/{signals[0]['signal_id']}/hypotheses")
    assert response.status_code == 503
    assert response.json()["detail"]["code"] == "reasoner_unavailable"
    assert client.get(f"/diagnostics/signals/{signals[0]['signal_id']}/hypotheses").json() == []
    assert client.post("/designs/diagnoses/any").json()["detail"]["code"] == "diagnosis_not_found"
    # Nothing from the synthetic demo path, and no identity columns, reach any real-mode response.
    for path in ("/diagnostics/signals", f"/diagnostics/signals/{signals[0]['signal_id']}/review-evidence"):
        body = json.dumps(client.get(path).json())
        assert not any(marker in body for marker in
                       ("demo_eval_", "synthetic-demo", "m4-demo-fixture", "Synthetic Agent"))
        assert not {"agent_name", "qa_name", "team_leader"} & set(body.split('"'))


def test_mode_endpoint_reports_installed_providers_not_the_label(app_state):
    app_state.state.demo_mode = "synthetic_demo"
    app_state.state.diagnostics = DiagnosticService([], UnavailableReasoner())
    unavailable = UnavailableDesignProvider()
    app_state.state.designs = DesignService(app_state.state.diagnostics, unavailable, unavailable)
    mode = TestClient(app_state).get("/diagnostics/mode").json()
    assert {key: mode[key] for key in mode if key not in ("intervention_provider", "solution_validator")} == {
        "mode": "synthetic_demo", "diagnostic_provider": "unavailable", "remote_diagnosis": "unavailable",
        "design_provider": "unavailable", "evaluation_count": 0, "signal_count": 0}
    fixture = DemoDesignFixture("sig_none")
    app_state.state.diagnostics = DiagnosticService([], ControlledTestReasoner({}))
    app_state.state.designs = DesignService(app_state.state.diagnostics, fixture, fixture,
                                            controlled_fixture=True)
    mode = TestClient(app_state).get("/diagnostics/mode").json()
    assert {key: mode[key] for key in mode if key not in ("intervention_provider", "solution_validator")} == {
        "mode": "synthetic_demo", "diagnostic_provider": "controlled_fixture",
        "remote_diagnosis": "local_fixture",
        "design_provider": "controlled_fixture", "evaluation_count": 0, "signal_count": 0}

    class SomeProvider:
        async def diagnose(self, bundle):
            return {}

        async def decide(self, context):
            return {}

        async def design(self, context, decision):
            return {}

    app_state.state.diagnostics = DiagnosticService([], SomeProvider())
    assert TestClient(app_state).get("/diagnostics/mode").json()["diagnostic_provider"] == "provider"
    # A real provider beside a fixture or an unavailable partner is still a provider.
    app_state.state.designs = DesignService(app_state.state.diagnostics, SomeProvider(), fixture)
    assert TestClient(app_state).get("/diagnostics/mode").json()["design_provider"] == "provider"
    app_state.state.designs = DesignService(app_state.state.diagnostics, unavailable, SomeProvider())
    assert TestClient(app_state).get("/diagnostics/mode").json()["design_provider"] == "provider"


def test_design_fixture_label_cannot_disagree_with_installed_providers():
    """`generation_mode` and `/diagnostics/mode` derive from the objects, not a constructor flag."""
    diagnostics = DiagnosticService([], UnavailableReasoner())
    unavailable = UnavailableDesignProvider()
    fixture = DemoDesignFixture("sig_none")

    class RemoteLookingProvider:
        async def decide(self, context):
            return {}

        async def design(self, context, decision):
            return {}

    for intervention, training in ((unavailable, unavailable),
                                   (RemoteLookingProvider(), RemoteLookingProvider()),
                                   (RemoteLookingProvider(), fixture)):
        with pytest.raises(DesignError, match="controlled_fixture flag disagrees") as refused:
            DesignService(diagnostics, intervention, training, controlled_fixture=True)
        assert refused.value.code == "provider_mismatch"
    with pytest.raises(DesignError, match="controlled_fixture flag disagrees"):
        DesignService(diagnostics, fixture, fixture)  # A fixture may not pose as a provider either.
    assert DesignService(diagnostics, fixture, fixture, controlled_fixture=True).controlled_fixture is True
    assert DesignService(diagnostics, fixture, unavailable, controlled_fixture=True).controlled_fixture is True
    assert DesignService(diagnostics, unavailable, unavailable).controlled_fixture is False


def test_default_app_mode_is_unconfigured_with_no_providers():
    assert app.state.demo_mode == "unconfigured"
    assert TestClient(app).get("/diagnostics/mode").json() == {
        "mode": "unconfigured", "diagnostic_provider": "unavailable", "remote_diagnosis": "unavailable",
        "design_provider": "unavailable", "intervention_provider": "unavailable",
        "solution_validator": "unavailable", "evaluation_count": 0, "signal_count": 0}


def test_confidential_paths_are_ignored_and_not_tracked():
    root = Path(__file__).resolve().parents[3]
    paths = [f"data/raw/{name}" for name in demo.REQUIRED_WORKBOOKS.values()]
    paths.append("data/processed/evaluations.jsonl")
    paths.append("data/raw/archive/anything.xlsx")
    for path in paths:
        assert subprocess.run(["git", "check-ignore", "-q", path], cwd=root).returncode == 0
        assert subprocess.run(["git", "ls-files", "--error-unmatch", path], cwd=root,
                              capture_output=True).returncode != 0
    tracked = subprocess.run(["git", "ls-files", "data"], cwd=root, capture_output=True, text=True)
    assert set(tracked.stdout.split()) == {"data/raw/.gitkeep", "data/processed/.gitkeep"}


def test_explicit_script_fails_safely_without_workbooks(tmp_path):
    root = Path(__file__).resolve().parents[3]
    result = subprocess.run([sys.executable, "scripts/run_results_cx_demo.py", "--input", str(tmp_path)],
                            cwd=root, capture_output=True, text=True)
    assert result.returncode == 2
    assert "missing_demo_workbooks" in result.stderr
    assert "Mode: REAL RESULTS CX" not in result.stdout
    assert "Traceback" not in result.stderr


def test_explicit_script_refuses_configured_evaluations_path(tmp_path):
    root = Path(__file__).resolve().parents[3]
    result = subprocess.run([sys.executable, "scripts/run_results_cx_demo.py", "--input", str(tmp_path)],
                            cwd=root, capture_output=True, text=True,
                            env={"PATH": "", "COACHLENS_API_DIAGNOSTIC_EVALUATIONS_PATH": str(tmp_path / "x.jsonl")})
    assert result.returncode == 2
    assert "refuses COACHLENS_API_DIAGNOSTIC_EVALUATIONS_PATH" in result.stderr
