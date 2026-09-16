"""Synthetic workbook tests: no ResultsCX names, comments, or identifiers."""

from datetime import datetime
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
import pytest

from app.results_cx import (PipelineValidationError, analyze, discover_sources,
                            normalize, profile_sources, write_jsonl)
from app.results_cx.models import Domain, Evaluation
from app.results_cx.pipeline import evaluation_id


COMMON = ["QA Name", "Team Leader", "AgentNames", "Date of Call", "Questions",
          "Answers", "Feedback answers", "MaxScore", "Score"]
MARKERS = {"business_process": "Business Pass Count", "compliance": "Compliance Pass Count",
           "member_experience": "QA Pass Count"}


def make_workbook(root: Path, domain: str, *, rows=None, filename=None, headers=None):
    root.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Export"
    columns = headers or (COMMON[:4] + (["Date of Evaluation"] if domain == "compliance" else []) + COMMON[4:] + [MARKERS[domain]])
    sheet.append(columns)
    if rows is None:
        rows = [
            {"QA Name": "Evaluator A", "Team Leader": "Lead A", "AgentNames": agent,
             "Date of Call": datetime(2026, 1, day), "Date of Evaluation": datetime(2026, 1, day + 1),
             "Questions": question, "Answers": answer, "Feedback answers": feedback,
             "MaxScore": 10, "Score": 10 if answer.casefold() == "yes" else 0,
             MARKERS[domain]: 1 if answer.casefold() == "yes" else 0}
            for agent, day in (("Agent A", 1), ("Agent B", 2))
            for question, answer, feedback in (("Criterion A", "yes", None), ("Criterion B", "No", "Synthetic note"))
        ]
    for row in rows:
        sheet.append([row.get(column) for column in columns])
    sheet.append(["Total" if column == "QA Name" else None for column in columns])
    path = root / (filename or f"renamed-{domain}.xlsx")
    workbook.save(path)
    return path


@pytest.fixture
def sources(tmp_path):
    root = tmp_path / "raw"
    for domain in MARKERS:
        make_workbook(root, domain)
    return discover_sources(root)


def test_schema_discovery_and_profile(sources):
    assert set(sources) == set(Domain)
    assert all(source.sheet == "Export" for source in sources.values())
    profiles = profile_sources(sources)
    assert len(profiles) == 3
    assert all(profile["criterion_row_count"] == 4 for profile in profiles)
    assert all(profile["distinct_evaluated_calls"] == 2 for profile in profiles)
    assert all(profile["agent_count"] == 2 for profile in profiles)
    assert all(profile["answer_distribution"] == {"Yes": 2, "No": 2} for profile in profiles)
    assert all("Agent A" not in json.dumps(profile) and "Synthetic note" not in json.dumps(profile) for profile in profiles)


def test_discovery_uses_schema_not_sheet_or_filename(tmp_path):
    root = tmp_path / "raw"
    paths = [make_workbook(root, domain, filename=f"source-{index}.xlsx")
             for index, domain in enumerate(MARKERS)]
    workbook = load_workbook(paths[0])
    workbook["Export"].title = "QA Rows"
    workbook.save(paths[0])
    found = discover_sources(root)
    assert found[Domain.BUSINESS_PROCESS].sheet == "QA Rows"


def test_original_answer_and_question_are_preserved(sources):
    source = sources[Domain.BUSINESS_PROCESS]
    workbook = load_workbook(source.path)
    sheet = workbook["Export"]
    sheet.cell(2, 5).value = " Criterion A "
    sheet.cell(2, 6).value = " Yes "
    workbook.save(source.path)
    result = normalize(sources)[0].criteria[0]
    assert result.question == " Criterion A "
    assert result.answer == " Yes "
    assert result.passed is True


def test_missing_and_duplicate_domains(tmp_path):
    root = tmp_path / "raw"
    make_workbook(root, "business_process")
    with pytest.raises(PipelineValidationError, match="missing_domain"):
        discover_sources(root)
    make_workbook(root, "compliance")
    make_workbook(root, "member_experience")
    make_workbook(root, "compliance", filename="another-compliance.xlsx")
    with pytest.raises(PipelineValidationError, match="duplicate_domain"):
        discover_sources(root)


def test_rejects_unsupported_and_ambiguous_schema(tmp_path):
    root = tmp_path / "raw"
    for domain in MARKERS:
        make_workbook(root, domain)
    make_workbook(root, "business_process", filename="bad.xlsx", headers=["Wrong", MARKERS["business_process"]])
    with pytest.raises(PipelineValidationError, match="missing_columns"):
        discover_sources(root)
    (root / "bad.xlsx").unlink()
    make_workbook(root, "business_process", filename="bad.xlsx", headers=COMMON + list(MARKERS.values()))
    with pytest.raises(PipelineValidationError, match="ambiguous_domain"):
        discover_sources(root)


def test_normalization_join_id_and_lineage(sources):
    evaluations = normalize(sources)
    assert len(evaluations) == 2
    assert all(isinstance(e, Evaluation) and len(e.criteria) == 6 for e in evaluations)
    assert all({c.domain for c in e.criteria} == set(Domain) for e in evaluations)
    first = evaluations[0]
    assert first.internal_id == evaluation_id((first.agent_name, first.call_date, first.qa_name, first.team_leader))
    assert first.internal_id == normalize(sources)[0].internal_id
    assert first.evaluation_date.isoformat() == "2026-01-02"
    assert {c.lineage.source_sheet for c in first.criteria} == {"Export"}
    assert {c.lineage.excel_row for c in first.criteria} == {2, 3}
    assert {c.lineage.source_filename for c in first.criteria} == {s.path.name for s in sources.values()}
    assert all(c.passed == (c.answer.casefold() == "yes") for c in first.criteria)


def test_analytics_and_output_is_local_only(sources, tmp_path):
    evaluations = normalize(sources)
    stats = analyze(evaluations)
    domain_stats = [s for s in stats if s.question is None]
    assert len(domain_stats) == 3
    for stat in domain_stats:
        assert (stat.evaluations, stat.pass_count, stat.fail_count, stat.feedback_count) == (2, 2, 2, 2)
        assert (stat.pass_rate, stat.fail_rate, stat.feedback_coverage) == (0.5, 0.5, 0.5)
        assert (stat.max_score_total, stat.attained_score_total, stat.score_rate) == (40, 20, 0.5)
        assert len(stat.lineages) == 4
        assert "lineages" not in stat.model_dump()
    criterion = next(s for s in stats if s.domain == Domain.COMPLIANCE and s.question == "Criterion A")
    assert (criterion.evaluations, criterion.pass_count, criterion.fail_count, criterion.feedback_count) == (2, 2, 0, 0)
    target = write_jsonl(evaluations, tmp_path / "data" / "processed")
    loaded = [Evaluation.model_validate_json(line) for line in target.read_text().splitlines()]
    assert [e.internal_id for e in loaded] == [e.internal_id for e in evaluations]
    with pytest.raises(PipelineValidationError, match="unsafe_output"):
        write_jsonl(evaluations, tmp_path / "export")


def test_duplicate_candidate_key_fails(sources):
    source = sources[Domain.COMPLIANCE]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Export"
    rows = list(load_workbook(source.path, data_only=True)["Export"].values)
    for row in rows:
        sheet.append(row)
    sheet.append(rows[1])
    workbook.save(source.path)
    with pytest.raises(PipelineValidationError, match="duplicate_candidate_key"):
        normalize(sources)


def test_missing_counterpart_and_metadata_mismatch(sources):
    source = sources[Domain.MEMBER_EXPERIENCE]
    workbook = load_workbook(source.path)
    sheet = workbook["Export"]
    sheet.delete_rows(2, 2)
    workbook.save(source.path)
    with pytest.raises(PipelineValidationError, match="missing_counterpart"):
        normalize(sources)
    # Restore the fixture and change a join metadata field in both criterion rows.
    make_workbook(source.path.parent, "member_experience")
    workbook = load_workbook(source.path)
    sheet = workbook["Export"]
    for row in (2, 3):
        sheet.cell(row, 1).value = "Evaluator B"
    workbook.save(source.path)
    with pytest.raises(PipelineValidationError, match="share agent/date with different evaluator"):
        normalize(sources)


def test_invalid_answer_and_conflicting_evaluation_date(sources):
    source = sources[Domain.BUSINESS_PROCESS]
    workbook = load_workbook(source.path)
    workbook["Export"].cell(2, 6).value = "Maybe"
    workbook.save(source.path)
    with pytest.raises(PipelineValidationError, match="invalid_answer"):
        normalize(sources)
    make_workbook(source.path.parent, "business_process")
    source = sources[Domain.COMPLIANCE]
    workbook = load_workbook(source.path)
    workbook["Export"].cell(3, 5).value = datetime(2026, 1, 10)
    workbook.save(source.path)
    with pytest.raises(PipelineValidationError, match="inconsistent_metadata"):
        normalize(sources)


# --- Review additions: failure behaviour, lineage stability, determinism, privacy ---

COMPLIANCE_COLUMNS = COMMON[:4] + ["Date of Evaluation"] + COMMON[4:] + [MARKERS["compliance"]]


def set_cells(path: Path, edits: dict[tuple[int, int], object]):
    """Assign (row, column) -> value; unlike Worksheet.cell(), None is assigned too."""
    workbook = load_workbook(path)
    sheet = workbook["Export"]
    for (row, column), value in edits.items():
        sheet.cell(row, column).value = value
    workbook.save(path)


@pytest.mark.parametrize("column, value, code", [
    ("AgentNames", None, "missing_field"),
    ("AgentNames", "   ", "missing_field"),
    ("AgentNames", 12345, "missing_field"),
    ("QA Name", None, "missing_field"),
    ("Team Leader", None, "missing_field"),
    ("Questions", None, "missing_field"),
    ("Answers", None, "missing_field"),
    ("Date of Call", None, "invalid_date"),
    ("Date of Call", "01/02/2026", "invalid_date"),
    ("Date of Call", datetime(2026, 1, 1, 9, 30), "timestamp_granularity"),
    ("Score", None, "invalid_score"),
    ("Score", "n/a-sentinel", "invalid_score"),
    ("Score", -1, "invalid_score"),
    ("Score", True, "invalid_score"),
    ("Score", 11, "score_exceeds_max"),
    ("MaxScore", None, "invalid_score"),
    ("Feedback answers", 42, "invalid_feedback"),
    (MARKERS["compliance"], 2, "invalid_marker"),
])
def test_malformed_criterion_rows_fail_without_leaking_values(sources, column, value, code):
    path = sources[Domain.COMPLIANCE].path
    set_cells(path, {(2, COMPLIANCE_COLUMNS.index(column) + 1): value})
    with pytest.raises(PipelineValidationError, match=code) as info:
        normalize(sources)
    message = str(info.value)
    assert "row 2" in message and "Agent A" not in message
    if isinstance(value, str) and value.strip():
        assert value not in message


def test_missing_score_is_not_treated_as_zero(sources):
    path = sources[Domain.COMPLIANCE].path
    set_cells(path, {(3, COMPLIANCE_COLUMNS.index("Score") + 1): None})
    with pytest.raises(PipelineValidationError, match="invalid_score"):
        normalize(sources)


def test_answer_must_agree_with_pass_marker(sources):
    path = sources[Domain.COMPLIANCE].path
    marker = COMPLIANCE_COLUMNS.index(MARKERS["compliance"]) + 1
    set_cells(path, {(2, marker): None})  # row 2 answered yes
    with pytest.raises(PipelineValidationError, match="answer_marker_conflict"):
        normalize(sources)
    set_cells(path, {(2, marker): 1, (3, marker): 1})  # row 3 answered No
    with pytest.raises(PipelineValidationError, match="answer_marker_conflict"):
        normalize(sources)
    set_cells(path, {(3, marker): 0})  # explicit zero is an acceptable "not passed" marker
    assert len(normalize(sources)) == 2


def test_different_team_leader_fails_to_join(sources):
    path = sources[Domain.MEMBER_EXPERIENCE].path
    set_cells(path, {(2, 2): "Lead B", (3, 2): "Lead B"})
    with pytest.raises(PipelineValidationError, match="1 share agent/date with different evaluator or leader"):
        normalize(sources)


def test_identity_trims_whitespace_but_not_casing(sources):
    path = sources[Domain.COMPLIANCE].path
    set_cells(path, {(2, 3): "  Agent A ", (3, 3): "Agent A\t"})
    evaluations = normalize(sources)
    assert evaluations[0].agent_name == "Agent A"
    assert evaluations[0].internal_id == evaluation_id(("Agent A", evaluations[0].call_date, "Evaluator A", "Lead A"))
    set_cells(path, {(2, 3): "agent a", (3, 3): "agent a"})
    with pytest.raises(PipelineValidationError, match="missing_counterpart"):
        normalize(sources)


def test_second_call_with_same_identity_is_rejected(sources):
    # A second call by the same agent on the same date, evaluated by the same people, repeats
    # every criterion in every domain; the pipeline must refuse rather than merge the calls.
    for source in sources.values():
        workbook = load_workbook(source.path)
        sheet = workbook["Export"]
        duplicates = [[sheet.cell(row, column).value for column in range(1, sheet.max_column + 1)] for row in (2, 3)]
        sheet.delete_rows(sheet.max_row)
        for values in duplicates:
            sheet.append(values)
        workbook.save(source.path)
    with pytest.raises(PipelineValidationError, match="duplicate_candidate_key"):
        normalize(sources)


def test_incomplete_evaluation_is_visible_in_per_criterion_counts(sources):
    workbook = load_workbook(sources[Domain.COMPLIANCE].path)
    workbook["Export"].delete_rows(3)
    workbook.save(sources[Domain.COMPLIANCE].path)
    evaluations = normalize(sources)
    assert [len(e.criteria) for e in evaluations] == [5, 6]
    stats = {s.question: s for s in analyze(evaluations) if s.domain == Domain.COMPLIANCE}
    assert (stats[None].evaluations, stats[None].pass_count + stats[None].fail_count) == (2, 3)
    assert (stats["Criterion B"].evaluations, stats["Criterion B"].fail_count) == (1, 1)
    profile = next(p for p in profile_sources(sources) if p["domain"] == "compliance")
    assert profile["criterion_row_counts"] == {"Criterion A": 2, "Criterion B": 1}
    assert profile["criteria_per_call"] == {1: 1, 2: 1}


def test_rows_with_evidence_but_no_identity_are_not_silently_dropped(sources):
    path = sources[Domain.COMPLIANCE].path
    blank = {(2, column): None for column in range(1, len(COMPLIANCE_COLUMNS) + 1)}
    set_cells(path, {**blank, (2, COMPLIANCE_COLUMNS.index("Feedback answers") + 1): "orphan comment"})
    with pytest.raises(PipelineValidationError, match="missing_field") as info:
        normalize(sources)
    assert "orphan comment" not in str(info.value)


def test_footer_totals_and_blank_rows_are_excluded(sources):
    path = sources[Domain.COMPLIANCE].path
    workbook = load_workbook(path)
    sheet = workbook["Export"]
    sheet.append([None] * len(COMPLIANCE_COLUMNS))
    sheet.append(["Total" if c == "QA Name" else 40 if c == "MaxScore" else 30 if c == "Score" else 3 if c.endswith("Pass Count") else None
                  for c in COMPLIANCE_COLUMNS])
    sheet.append(["Applied filters:\nSynthetic note" if c == "QA Name" else None for c in COMPLIANCE_COLUMNS])
    workbook.save(path)
    profile = next(p for p in profile_sources(sources) if p["domain"] == "compliance")
    assert (profile["row_count"], profile["criterion_row_count"], profile["excluded_footer_or_blank_rows"]) == (8, 4, 4)
    assert profile["unique_counts"]["QA Name"] == 1 and profile["non_null_counts"]["Score"] == 4
    evaluations = normalize(sources)
    stats = next(s for s in analyze(evaluations) if s.domain == Domain.COMPLIANCE and s.question is None)
    assert (stats.pass_count + stats.fail_count, stats.max_score_total, stats.attained_score_total) == (4, 40, 20)


def test_lineage_survives_blank_rows_and_matches_openpyxl_coordinates(sources):
    path = sources[Domain.COMPLIANCE].path
    workbook = load_workbook(path)
    workbook["Export"].insert_rows(3)
    workbook["Export"].insert_rows(2)
    workbook.save(path)
    evaluations = normalize(sources)
    compliance = [c for e in evaluations for c in e.criteria if c.domain == Domain.COMPLIANCE]
    assert sorted(c.lineage.excel_row for c in compliance) == [3, 5, 6, 7]
    sheet = load_workbook(path)["Export"]
    header = [cell.value for cell in sheet[1]]
    for criterion in compliance:
        row = sheet[criterion.lineage.excel_row]
        assert row[header.index("Questions")].value == criterion.question
        assert row[header.index("Answers")].value == criterion.answer
        assert str(row[header.index("Score")].value) == str(criterion.attained_score)


def test_output_is_independent_of_row_order_and_filenames(tmp_path):
    def build(root: Path, shuffle: bool):
        for index, domain in enumerate(MARKERS):
            path = make_workbook(root, domain, filename=f"{'shuffled' if shuffle else 'plain'}-{index}.xlsx")
            if shuffle:
                workbook = load_workbook(path)
                sheet = workbook["Export"]
                rows = list(sheet.values)
                body = rows[1:-1]
                body.reverse()
                sheet.delete_rows(1, sheet.max_row)
                for values in [rows[0], *body, rows[-1]]:
                    sheet.append(values)
                workbook.save(path)
        return normalize(discover_sources(root))

    plain, shuffled = build(tmp_path / "a", False), build(tmp_path / "b", True)
    assert [e.internal_id for e in plain] == [e.internal_id for e in shuffled]
    canonical = lambda evaluations: [  # noqa: E731
        {**e.model_dump(exclude={"criteria"}),
         "criteria": sorted((c.model_dump(exclude={"lineage"}) for c in e.criteria), key=lambda c: (c["domain"], c["question"]))}
        for e in evaluations]
    assert canonical(plain) == canonical(shuffled)
    assert {c.lineage.excel_row for c in plain[0].criteria} != {c.lineage.excel_row for c in shuffled[0].criteria}


def test_jsonl_round_trip_preserves_decimals_dates_and_unicode(tmp_path):
    root = tmp_path / "raw"
    for domain in MARKERS:
        path = make_workbook(root, domain)
        edits = {(2, 3): "Ágent Ñ 日本", (3, 3): "Ágent Ñ 日本", (2, 4): datetime(2026, 1, 1)}
        if domain == "compliance":
            edits[(2, 10)] = 11.11
            edits[(2, 9)] = 11.11
        set_cells(path, edits)
    evaluations = normalize(discover_sources(root))
    target = write_jsonl(evaluations, tmp_path / "data" / "processed")
    first = json.loads(target.read_text(encoding="utf-8").splitlines()[0])
    assert first["agent_name"] == "Ágent Ñ 日本" and first["call_date"] == "2026-01-01"
    compliance = next(c for c in first["criteria"] if c["domain"] == "compliance" and c["question"] == "Criterion A")
    assert (compliance["attained_score"], compliance["max_score"]) == ("11.11", "11.11")
    assert Evaluation.model_validate_json(target.read_text(encoding="utf-8").splitlines()[0]) == evaluations[0]
    assert target.read_bytes() == b"".join(e.model_dump_json().encode("utf-8") + b"\n" for e in evaluations)


def test_write_jsonl_never_leaves_partial_output(sources, tmp_path):
    evaluations = normalize(sources)
    output = tmp_path / "data" / "processed"
    target = write_jsonl(evaluations, output)
    before = target.read_bytes()

    def failing():
        yield evaluations[0]
        raise RuntimeError("simulated failure")

    with pytest.raises(RuntimeError):
        write_jsonl(failing(), output)
    assert target.read_bytes() == before
    assert sorted(p.name for p in output.iterdir()) == ["evaluations.jsonl"]


def test_workbook_without_dimension_tag_parses(sources):
    import re
    import zipfile

    path = sources[Domain.COMPLIANCE].path
    rewritten = path.with_suffix(".rewrite")
    with zipfile.ZipFile(path) as source_zip, zipfile.ZipFile(rewritten, "w", zipfile.ZIP_DEFLATED) as target_zip:
        for item in source_zip.infolist():
            data = source_zip.read(item.filename)
            if item.filename.startswith("xl/worksheets/sheet"):
                data = re.sub(rb"<dimension[^>]*/>", b"", data)
            target_zip.writestr(item, data)
    rewritten.replace(path)
    assert load_workbook(path, read_only=True)["Export"].max_row is None
    rediscovered = discover_sources(path.parent)
    assert rediscovered[Domain.COMPLIANCE].sheet == "Export"
    evaluations = normalize(rediscovered)
    assert len(evaluations) == 2 and all(len(e.criteria) == 6 for e in evaluations)


def test_analytics_bucket_criteria_by_trimmed_question(sources):
    set_cells(sources[Domain.COMPLIANCE].path, {(2, 6): " Criterion A "})
    evaluations = normalize(sources)
    compliance = [s for s in analyze(evaluations) if s.domain == Domain.COMPLIANCE]
    assert sorted(s.question or "" for s in compliance) == ["", "Criterion A", "Criterion B"]
    assert next(s for s in compliance if s.question == "Criterion A").evaluations == 2
    # The original cell text is still preserved on the criterion record.
    assert any(c.question == " Criterion A " for c in evaluations[0].criteria)


def test_analytics_rates_use_criterion_row_denominators(sources):
    stats = analyze(normalize(sources))
    for stat in stats:
        total = stat.pass_count + stat.fail_count
        assert total == len(stat.lineages)
        assert stat.pass_rate + stat.fail_rate == 1
        assert stat.pass_rate * total == stat.pass_count
        assert stat.feedback_coverage * total == stat.feedback_count
        assert stat.evaluations <= total
        assert stat.score_rate * stat.max_score_total == stat.attained_score_total
