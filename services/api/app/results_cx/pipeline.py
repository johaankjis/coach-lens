"""Strict, reproducible parsing of long-form ResultsCX QA exports.

No model inference, statistical heuristics, or confidential output logging occurs here.
"""

from collections import Counter, defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from hashlib import sha256
import json
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

from .models import CriterionResult, Domain, Evaluation, SourceLineage, Statistic


COMMON = {"QA Name", "Team Leader", "AgentNames", "Date of Call", "Questions",
          "Answers", "Feedback answers", "MaxScore", "Score"}
MARKERS = {Domain.BUSINESS_PROCESS: "Business Pass Count",
           Domain.COMPLIANCE: "Compliance Pass Count",
           Domain.MEMBER_EXPERIENCE: "QA Pass Count"}
OPTIONAL = {"Date of Evaluation"}
# Export footers ("Total" and the applied-filter note) only populate these columns.
FOOTER_ALLOWED = {"QA Name", "MaxScore", "Score"} | set(MARKERS.values())


class PipelineValidationError(ValueError):
    """Actionable aggregate validation findings, without source values."""

    def __init__(self, code: str, message: str):
        self.code = code
        super().__init__(f"{code}: {message}")


class Source:
    def __init__(self, path: Path, sheet: str, headers: tuple[str, ...], domain: Domain):
        self.path, self.sheet, self.headers, self.domain = path, sheet, headers, domain


def _header(ws) -> tuple[str, ...] | None:
    # Read-only sheets report max_row as None when the XML has no <dimension>, so probe the row itself.
    values = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if not values or any(not isinstance(v, str) or not v.strip() for v in values):
        return None
    headers = tuple(v.strip() for v in values)
    if len(set(headers)) != len(headers):
        raise PipelineValidationError("duplicate_columns", f"{ws.title}: repeated header names")
    return headers


def _classify(headers: tuple[str, ...]) -> Domain | None:
    names = set(headers)
    matches = [domain for domain, marker in MARKERS.items() if marker in names]
    if matches and not COMMON <= names:
        raise PipelineValidationError("missing_columns", f"QA sheet is missing required columns: {', '.join(sorted(COMMON - names))}")
    if not COMMON <= names:
        return None
    if len(matches) != 1:
        raise PipelineValidationError("ambiguous_domain", "QA sheet has zero or multiple domain marker columns")
    allowed = COMMON | OPTIONAL | set(MARKERS.values())
    extra = names - allowed
    if extra:
        raise PipelineValidationError("unsupported_columns", f"Unexpected columns: {', '.join(sorted(extra))}")
    return matches[0]


def discover_sources(input_dir: Path | str) -> dict[Domain, Source]:
    root = Path(input_dir)
    if not root.is_dir():
        raise PipelineValidationError("missing_input", f"Input directory does not exist: {root}")
    paths = sorted(p for p in root.rglob("*.xlsx") if not p.name.startswith("~$"))
    if not paths:
        raise PipelineValidationError("no_workbooks", f"No .xlsx workbooks found under {root}")
    found: dict[Domain, Source] = {}
    for path in paths:
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
        except Exception as exc:
            raise PipelineValidationError("unreadable_workbook", f"Cannot read {path.name}: {exc.__class__.__name__}") from exc
        try:
            candidates = []
            for ws in wb:
                headers = _header(ws)
                if headers is None:
                    continue
                domain = _classify(headers)
                if domain is not None:
                    candidates.append(Source(path, ws.title, headers, domain))
            if len(candidates) != 1:
                raise PipelineValidationError("unsupported_structure", f"{path.name}: expected exactly one long-form QA sheet; found {len(candidates)}")
            source = candidates[0]
            if source.domain in found:
                raise PipelineValidationError("duplicate_domain", f"{source.domain.value}: {found[source.domain].path.name} and {path.name}")
            found[source.domain] = source
        finally:
            wb.close()
    missing = set(Domain) - set(found)
    if missing:
        raise PipelineValidationError("missing_domain", f"Missing required source domain(s): {', '.join(sorted(d.value for d in missing))}")
    return found


def _read_rows(source: Source):
    wb = load_workbook(source.path, read_only=True, data_only=True)
    try:
        ws = wb[source.sheet]
        for number, values in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
            if len(values) > len(source.headers) and any(v is not None for v in values[len(source.headers):]):
                raise PipelineValidationError("unsupported_structure", f"{source.path.name} row {number}: values beyond the header columns")
            row = dict(zip(source.headers, values))
            # Export totals, notes, and empty rows carry no identity, question, answer, or feedback.
            # Any other populated column means evidence, so the row must validate as a criterion row.
            eligible = any(row.get(k) is not None for k in source.headers if k not in FOOTER_ALLOWED)
            yield number, row, eligible
    finally:
        wb.close()


def _date(value, source: Source, number: int, field: str) -> date:
    if isinstance(value, datetime):
        if value.time().isoformat() != "00:00:00":
            raise PipelineValidationError("timestamp_granularity", f"{source.path.name} row {number}: {field} contains a time; date-only identity would be ambiguous")
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        try:
            return date.fromisoformat(value.strip())
        except ValueError:
            pass
    raise PipelineValidationError("invalid_date", f"{source.path.name} row {number}: {field} must be an Excel date or ISO date")


def _text(value, source: Source, number: int, field: str) -> str:
    if not isinstance(value, str) or not value.strip():
        raise PipelineValidationError("missing_field", f"{source.path.name} row {number}: {field} must be nonempty text")
    return value.strip()


def _score(value, source: Source, number: int, field: str) -> Decimal:
    if value is None or isinstance(value, bool):
        raise PipelineValidationError("invalid_score", f"{source.path.name} row {number}: {field} must be numeric")
    try:
        score = Decimal(str(value))
    except (InvalidOperation, ValueError):
        raise PipelineValidationError("invalid_score", f"{source.path.name} row {number}: {field} must be numeric") from None
    if not score.is_finite() or score < 0:
        raise PipelineValidationError("invalid_score", f"{source.path.name} row {number}: {field} must be finite and nonnegative")
    return score


def _parsed(source: Source):
    for number, row, eligible in _read_rows(source):
        if not eligible:
            continue
        agent = _text(row.get("AgentNames"), source, number, "AgentNames")
        call_date = _date(row.get("Date of Call"), source, number, "Date of Call")
        qa = _text(row.get("QA Name"), source, number, "QA Name")
        leader = _text(row.get("Team Leader"), source, number, "Team Leader")
        question = _text(row.get("Questions"), source, number, "Questions")
        answer = _text(row.get("Answers"), source, number, "Answers")
        if answer.casefold() not in {"yes", "no"}:
            raise PipelineValidationError("invalid_answer", f"{source.path.name} row {number}: Answers must be Yes or No")
        passed = answer.casefold() == "yes"
        # The export's own pass flag must agree with the answer; disagreement means a corrupted row.
        marker_name = MARKERS[source.domain]
        marker = row.get(marker_name)
        flag = Decimal(0) if marker is None else _score(marker, source, number, marker_name)
        if flag not in (0, 1):
            raise PipelineValidationError("invalid_marker", f"{source.path.name} row {number}: {marker_name} must be blank, 0, or 1")
        if (flag == 1) != passed:
            raise PipelineValidationError("answer_marker_conflict", f"{source.path.name} row {number}: Answers disagrees with {marker_name}")
        max_score = _score(row.get("MaxScore"), source, number, "MaxScore")
        score = _score(row.get("Score"), source, number, "Score")
        if score > max_score:
            raise PipelineValidationError("score_exceeds_max", f"{source.path.name} row {number}: Score exceeds MaxScore")
        feedback = row.get("Feedback answers")
        if feedback is not None and not isinstance(feedback, str):
            raise PipelineValidationError("invalid_feedback", f"{source.path.name} row {number}: Feedback answers must be text")
        evaluation_date = None
        if row.get("Date of Evaluation") is not None:
            evaluation_date = _date(row["Date of Evaluation"], source, number, "Date of Evaluation")
        key = (agent, call_date, qa, leader)
        result = CriterionResult(domain=source.domain, question=row["Questions"], answer=row["Answers"],
                                 passed=passed, max_score=max_score,
                                 attained_score=score, evaluator_feedback=feedback if feedback and feedback.strip() else None,
                                 lineage=SourceLineage(source_filename=source.path.name,
                                                       source_sheet=source.sheet, excel_row=number))
        yield key, evaluation_date, result


def evaluation_id(key: tuple[str, date, str, str]) -> str:
    """Full SHA-256 over versioned, ordered canonical key fields."""
    payload = json.dumps(["resultscx-v1", key[0], key[1].isoformat(), key[2], key[3]],
                         ensure_ascii=False, separators=(",", ":"))
    return "rcx1_" + sha256(payload.encode("utf-8")).hexdigest()


def normalize(sources: dict[Domain, Source]) -> list[Evaluation]:
    groups: dict[Domain, dict[tuple, list[CriterionResult]]] = {}
    dates: dict[tuple, set[date]] = defaultdict(set)
    for domain, source in sources.items():
        domain_groups: dict[tuple, list[CriterionResult]] = defaultdict(list)
        for key, evaluation_date, criterion in _parsed(source):
            domain_groups[key].append(criterion)
            if evaluation_date is not None:
                dates[key].add(evaluation_date)
        if not domain_groups:
            raise PipelineValidationError("empty_domain", f"{domain.value}: no criterion rows")
        for key, criteria in domain_groups.items():
            questions = [c.question.strip() for c in criteria]
            if len(questions) != len(set(questions)):
                row_numbers = [c.lineage.excel_row for c in criteria]
                raise PipelineValidationError("duplicate_candidate_key", f"{domain.value}: repeated criterion within one identity group, near rows {min(row_numbers)}-{max(row_numbers)}")
        groups[domain] = domain_groups
    all_keys = set.union(*(set(g) for g in groups.values()))
    for domain, domain_groups in groups.items():
        missing = all_keys - set(domain_groups)
        if missing:
            weaker = {(k[0], k[1]) for k in missing}
            observed = {(k[0], k[1]) for k in domain_groups}
            mismatched = len(weaker & observed)
            raise PipelineValidationError("missing_counterpart", f"{domain.value}: {len(missing)} unmatched identity group(s); {mismatched} share agent/date with different evaluator or leader metadata")
    inconsistent = {key for key, variants in dates.items() if len(variants) > 1}
    if inconsistent:
        raise PipelineValidationError("inconsistent_metadata", f"{len(inconsistent)} identity group(s) contain conflicting Date of Evaluation values")
    evaluations = []
    for key in sorted(all_keys, key=lambda k: (k[1], k[0], k[2], k[3])):
        criteria = [criterion for domain in Domain for criterion in groups[domain][key]]
        evaluations.append(Evaluation(internal_id=evaluation_id(key), agent_name=key[0],
                                      call_date=key[1], qa_name=key[2], team_leader=key[3],
                                      evaluation_date=next(iter(dates[key])) if dates[key] else None,
                                      criteria=criteria))
    if len({e.internal_id for e in evaluations}) != len(evaluations):
        raise PipelineValidationError("id_collision", "Distinct identity keys produced the same internal ID")
    return evaluations


def _rate(numerator: int, denominator: int) -> Decimal:
    return (Decimal(numerator) / Decimal(denominator)) if denominator else Decimal(0)


def analyze(evaluations: Iterable[Evaluation]) -> list[Statistic]:
    buckets: dict[tuple[Domain, str | None], list[CriterionResult]] = defaultdict(list)
    seen: dict[tuple[Domain, str | None], set[str]] = defaultdict(set)
    for evaluation in evaluations:
        for result in evaluation.criteria:
            for question in (None, result.question.strip()):
                bucket = (result.domain, question)
                buckets[bucket].append(result)
                seen[bucket].add(evaluation.internal_id)
    output = []
    for (domain, question), results in sorted(buckets.items(), key=lambda item: (item[0][0].value, item[0][1] or "")):
        passes = sum(r.passed for r in results)
        total = len(results)
        comments = sum(bool(r.evaluator_feedback) for r in results)
        maximum = sum((r.max_score for r in results), Decimal(0))
        attained = sum((r.attained_score for r in results), Decimal(0))
        output.append(Statistic(domain=domain, question=question, evaluations=len(seen[(domain, question)]),
                                pass_count=passes, fail_count=total-passes,
                                pass_rate=_rate(passes, total), fail_rate=_rate(total-passes, total),
                                feedback_count=comments, feedback_coverage=_rate(comments, total),
                                max_score_total=maximum, attained_score_total=attained,
                                score_rate=attained / maximum if maximum else Decimal(0),
                                lineages=[r.lineage for r in results]))
    return output


def profile_sources(sources: dict[Domain, Source]) -> list[dict]:
    profiles = []
    for domain, source in sources.items():
        wb = load_workbook(source.path, read_only=True, data_only=True)
        try:
            sheets = wb.sheetnames
        finally:
            wb.close()
        column_non_null = Counter()
        column_unique: dict[str, set] = defaultdict(set)
        physical = criterion_rows = footers = 0
        for _, row, eligible in _read_rows(source):
            physical += 1
            footers += not eligible
            criterion_rows += eligible
            if not eligible:
                continue
            for name, value in row.items():
                if value is not None and str(value).strip():
                    column_non_null[name] += 1
                    column_unique[name].add(str(value))
        parsed = list(_parsed(source))
        dates = [key[1] for key, _, _ in parsed]
        evaluation_dates = [evaluation_date for _, evaluation_date, _ in parsed if evaluation_date]
        profiles.append({"domain": domain.value, "source_filename": source.path.name,
                         "sheet_names": sheets, "selected_sheet": source.sheet,
                         "row_count": physical, "criterion_row_count": criterion_rows,
                         "excluded_footer_or_blank_rows": footers,
                         "column_names": list(source.headers),
                         "non_null_counts": {h: column_non_null[h] for h in source.headers},
                         "unique_counts": {h: len(column_unique[h]) for h in source.headers},
                         "call_date_range": [min(dates).isoformat(), max(dates).isoformat()] if dates else None,
                         "evaluation_date_range": [min(evaluation_dates).isoformat(), max(evaluation_dates).isoformat()] if evaluation_dates else None,
                         "agent_count": len({key[0] for key, _, _ in parsed}),
                         "distinct_evaluated_calls": len({key for key, _, _ in parsed}),
                         "criterion_names": sorted({result.question.strip() for _, _, result in parsed}),
                         "criterion_row_counts": dict(sorted(Counter(result.question.strip() for _, _, result in parsed).items())),
                         "criteria_per_call": dict(sorted(Counter(Counter(key for key, _, _ in parsed).values()).items())),
                         "answer_distribution": dict(Counter(result.answer.strip().casefold().capitalize() for _, _, result in parsed)),
                         "evaluator_feedback_coverage": str(_rate(sum(bool(result.evaluator_feedback) for _, _, result in parsed), len(parsed)))})
    return profiles


def write_jsonl(evaluations: Iterable[Evaluation], output_dir: Path | str) -> Path:
    output = Path(output_dir)
    if output.name != "processed" or output.resolve().parent.name != "data":
        raise PipelineValidationError("unsafe_output", "Output must be a data/processed directory")
    output.mkdir(parents=True, exist_ok=True)
    target = output / "evaluations.jsonl"
    temporary = output / "evaluations.jsonl.tmp"
    try:
        with temporary.open("w", encoding="utf-8") as stream:
            for evaluation in evaluations:
                stream.write(evaluation.model_dump_json() + "\n")
    except BaseException:
        temporary.unlink(missing_ok=True)
        raise
    temporary.replace(target)
    return target
